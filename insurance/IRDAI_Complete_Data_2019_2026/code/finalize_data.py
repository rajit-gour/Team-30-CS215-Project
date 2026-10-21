from pathlib import Path
from openpyxl import load_workbook
import pandas as pd
base=Path('/mnt/data/rajit_complete');out=base/'final_output';out.mkdir(exist_ok=True)
old=pd.read_csv(base/'complete_output/insurance_monthly_recovered.csv')
# directly read source figures: column in English table and the two immediately following Grand Total rows

def read_pair(code,contains,col=10):
 p=base/'raw_data'/f'IRDAI_{code}.xlsx'
 if not p.exists():p=base/'converted_xls'/f'IRDAI_{code}.xlsx'
 w=load_workbook(p,read_only=True,data_only=True)
 for s in w:
  if contains.lower() not in s.title.lower():continue
  rows=list(s.iter_rows(max_row=390,max_col=16,values_only=True))
  for i,r in enumerate(rows[:-2]):
   if str(r[1]).strip().lower()=='grand total':
    x,y=rows[i+1],rows[i+2]
    if 'individual single' in str(x[1]).lower() and 'individual non-single' in str(y[1]).lower():
     a,b=x[col],y[col]
     if isinstance(a,(float,int)) and isinstance(b,(float,int)):
      assert float(a).is_integer() and float(b).is_integer()
      return int(a),int(b),s.title,f'{chr(65+col)}{i+2};{chr(65+col)}{i+3}'
 raise Exception((code,contains,col))

fixes=[
 ('2019-03','070','2019',10,'English version appended beneath Hindi version'),
 ('2021-06','077','2021',10,'Incorrect printed header year For June 20201; workbook date and prior report match'),
 ('2021-07','087','2021',10,'Incorrect printed header year For July 20201; workbook date and prior report match'),
 ('2022-07','072','Sheet3',10,'Workbook mislabels July 2022 as June 2022: values agree with July 2023 previous-year comparison in IRDAI_033'),
 ('2022-09','013','Sept 2023',9,'September 2022 reconstructed from same-month prior-year comparison in September 2023 report'),
 ('2023-07','034','2024',9,'July 2023 reconstructed from same-month prior-year comparison in July 2024 report'),
]
fixlog=[]
for mm,code,sheet,col,reason in fixes:
 a,b,s,cells=read_pair(code,sheet,col)
 old=old[old.month!=mm]
 old=pd.concat([old,pd.DataFrame([{'month':mm,'individual_single':a,'individual_non_single':b,'individual_total':a+b,'source_file':f'IRDAI_{code}.xlsx','sheet':s,'status':'REVIEWED_RECOVERY'}])],ignore_index=True)
 fixlog.append(dict(month=mm,source=f'IRDAI_{code}.xlsx',sheet=s,cell_reference=cells,individual_single=a,individual_non_single=b,individual_total=a+b,method=reason))
# one duplicate conflict was actual July 2022 despite misleading June 2022 heading
jun=read_pair('050','2022',10)
old.loc[old.month=='2022-06',['individual_single','individual_non_single','individual_total','source_file','sheet','status']]=[jun[0],jun[1],jun[0]+jun[1],'IRDAI_050.xlsx',jun[2],'REVIEWED_DUPLICATE_CONFLICT']
fixlog.append(dict(month='2022-06',source='IRDAI_050.xlsx',sheet=jun[2],cell_reference=jun[3],individual_single=jun[0],individual_non_single=jun[1],individual_total=jun[0]+jun[1],method='Resolve conflict: IRDAI_072 contains July 2022 data despite headers labeled June; IRDAI_050 is actual June 2022'))
old=old.sort_values('month').reset_index(drop=True)
full=pd.period_range('2019-01','2026-09',freq='M').astype(str)
assert len(old)==93 and list(old.month)==list(full)
assert (old.individual_single+old.individual_non_single==old.individual_total).all()
# mark series as not independently audited
old['validation']='ARITHMETIC_CHECKED; YEAR/MONTH_CORRECTIONS_REVIEWED; NOT_INDEPENDENTLY_AUDITED'
old.to_csv(out/'insurance_monthly_2019_2026.csv',index=False)
pd.DataFrame(fixlog).to_csv(out/'recovery_and_correction_log.csv',index=False)
# source traceability
pd.read_csv(base/'source_file_map.csv').to_csv(out/'source_file_map.csv',index=False)
# detailed insurer policy extracts from identified current reports (91 candidates include old conflicting report; preserve provenance)
import shutil
shutil.copy(base/'complete_output/insurer_policy_detail.csv',out/'insurer_policy_detail_preliminary.csv')
shutil.copy(base/'complete_output/file_issues.csv',out/'original_parser_issues.csv')
# preserve other fields of earlier full workbook archive but note it was partial to detected sheets
shutil.copy(base/'output/all_reported_metrics.csv',out/'other_reported_metrics_preliminary.csv')
(out/'README.txt').write_text('''TEAM 30 | IRDAI LIFE INSURANCE | SOURCE-BASED MONTHLY EXTRACTION\n\nMain output: insurance_monthly_2019_2026.csv\n93 chronological months January 2019 through September 2026, with monthly new individual policy counts split into single and non-single premium. All observations originated from supplied IRDAI workbooks (including .xls).\n\nCorrections: recovery_and_correction_log.csv documents seven special cases. July 2022 is especially important: a workbook listed as July contains mislabeled June headers, but its numbers align with published July 2022 prior-year comparisons from a July 2023 workbook. The genuine June 2022 is taken from its own workbook.\n\nCAVEATS: This is a candidate research dataset, not a full independent audit. Recovered fields have been arithmetically checked; time-series and insurer subtotal validation are still necessary. The detailed insurer file and other reported metrics are preliminary and may contain reference information from incorrectly labeled reports, so use them only after checking source cell provenance.\n\nFiles: source_file_map.csv maps temporary identifiers back to exact original archive filenames.\n''',encoding='utf-8')
print('DONE',len(old),'month-range',old.month.min(),old.month.max(),'fixes',len(fixlog));print(pd.DataFrame(fixlog)[['month','individual_total','method']].to_string(index=False,max_colwidth=70))
